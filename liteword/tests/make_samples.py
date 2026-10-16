"""테스트용 샘플 문서(sample.docx)를 만들고 LibreOffice로 .doc/.rtf로 변환합니다.
사용법: python3 make_samples.py <출력 폴더>   (python-docx, Pillow, LibreOffice 필요)"""
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
im = Image.new('RGB', (200, 100), (40, 90, 200))
ImageDraw.Draw(im).text((70, 40), 'PNG', fill='white')
im.save(out / 'pic.png')

d = docx.Document()
d.add_heading('샘플 문서 Sample Document', 1)
p = d.add_paragraph('보통 문단: ')
p.add_run('굵게').bold = True
p.add_run(', ')
p.add_run('기울임').italic = True
p.add_run(', ')
p.add_run('밑줄').underline = True
p.add_run(', ')
r = p.add_run('빨간 14pt')
r.font.color.rgb = RGBColor(0xFF, 0, 0)
r.font.size = Pt(14)
p.add_run(', ')
r = p.add_run('바탕체')
r.font.name = 'Batang'
r._element.rPr.rFonts.set(qn('w:eastAsia'), '바탕')
p.add_run(', ')
p.add_run('형광펜').font.highlight_color = WD_COLOR_INDEX.YELLOW
p.add_run(' 끝.')
d.add_heading('목록 Lists', 2)
d.add_paragraph('글머리 하나', style='List Bullet')
d.add_paragraph('글머리 둘', style='List Bullet')
d.add_paragraph('들여쓴 글머리', style='List Bullet 2')
d.add_paragraph('번호 하나', style='List Number')
d.add_paragraph('번호 둘', style='List Number')
p = d.add_paragraph('가운데 정렬, 1.5줄 간격, 첫 줄 들여쓰기 문단입니다. ' * 3)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing = 1.5
p.paragraph_format.first_line_indent = Inches(0.5)
t = d.add_table(rows=3, cols=3)
t.style = 'Table Grid'
for i in range(3):
    for j in range(3):
        t.cell(i, j).text = f'칸{i}{j}'
t.cell(0, 0).merge(t.cell(0, 1)).text = '가로 병합'
t.cell(1, 2).merge(t.cell(2, 2)).text = '세로 병합'
p = d.add_paragraph('링크: ')
rid = d.part.relate_to('https://example.com/한글',
                       'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
h = OxmlElement('w:hyperlink')
h.set(qn('r:id'), rid)
nr = OxmlElement('w:r')
rpr = OxmlElement('w:rPr')
u = OxmlElement('w:u')
u.set(qn('w:val'), 'single')
rpr.append(u)
nr.append(rpr)
tt = OxmlElement('w:t')
tt.text = '예제 사이트'
nr.append(tt)
h.append(nr)
p._p.append(h)
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
d.add_paragraph('두 번째 쪽 그림:')
d.add_picture(str(out / 'pic.png'), width=Inches(2))
d.add_paragraph('Tab\t탭 문자와 English 텍스트 mixed 한글.')
d.save(out / 'sample.docx')
for fmt in ('doc', 'rtf'):
    subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', str(out), str(out / 'sample.docx')],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
print('samples in', out)
